"""Payroll Phases 6 and 7: audit log, unlock reason, bank file, PF ECR and ESI return."""
import csv
import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook

from appEmp.bank_file import HEADER as BANK_HEADER
from appEmp.bank_file import split_for_bank, write_bank_csv
from appEmp.ecr import ecr_line, split_for_ecr
from appEmp.esi_return import HEADER as ESI_HEADER
from appEmp.esi_return import build_esi_workbook, split_for_esi
from appEmp.models import EmployeeStatutoryProfile, PayrollAuditLog, PayrollRecord
from appEmp.test_payroll_service import MONTH, YEAR
from appEmp.test_payroll_views import PayrollViewBase

D = Decimal
HR_USERNAME = "zz_hr@example.com"


class FakeUser:
    def __init__(self, full="Zz Test", username="zz"):
        self.full = full
        self.username = username

    def get_full_name(self):
        return self.full

    def get_username(self):
        return self.username


def fake_employee(pk=1, **fields):
    data = dict(
        pk=pk, uuid="abcdef12-0000-0000-0000-000000000000", user=FakeUser(),
        bank_account_no="123456789012", ifsc_no="HDFC0001234", bank_name="HDFC",
        uan_no="100200300400", esic_no="1234567890",
    )
    data.update(fields)
    return SimpleNamespace(**data)


# ── Audit log and unlock reason (Phase 6) ───────────────────────────────
class AuditLogTests(PayrollViewBase):
    def locked_record(self, employee):
        return PayrollRecord.objects.get(employee=employee)

    def test_locking_writes_one_entry_with_the_locked_figures(self):
        a = self.full_employee()
        self.post("lock")
        entries = PayrollAuditLog.objects.filter(employee=a, action="LOCKED")
        self.assertEqual(entries.count(), 1)
        entry = entries.get()
        self.assertEqual((entry.year, entry.month), (YEAR, MONTH))
        self.assertEqual(entry.user_name, HR_USERNAME)
        self.assertEqual(entry.net_pay, self.locked_record(a).net_pay)
        self.assertEqual(entry.net_pay, D("58200.00"))
        self.assertEqual(entry.total_deductions, D("1800.00"))

    def test_locking_twice_and_calculating_add_no_entries(self):
        a = self.full_employee()
        self.post("lock")
        self.post("lock")
        self.post("calculate")
        self.assertEqual(PayrollAuditLog.objects.filter(employee=a).count(), 1)

    def test_unlock_needs_a_reason(self):
        a = self.full_employee()
        self.post("lock")
        for reason in ("", "   ", "too short", "x" * 256):
            response = self.post("unlock", reason=reason)
            self.assertTrue(self.locked_record(a).is_locked, reason)
            self.assertTrue(any("Give a reason" in m for m in self.messages(response)), reason)
        self.assertFalse(PayrollAuditLog.objects.filter(action="UNLOCKED").exists())

    def test_unlock_with_a_reason_is_logged(self):
        a = self.full_employee()
        self.post("lock")
        self.post("unlock", reason="Attendance was corrected for the 9th")
        self.assertFalse(self.locked_record(a).is_locked)
        entry = PayrollAuditLog.objects.get(employee=a, action="UNLOCKED")
        self.assertEqual(entry.reason, "Attendance was corrected for the 9th")
        self.assertEqual(entry.user_name, HR_USERNAME)

    def test_bulk_unlock_logs_every_record_with_the_same_reason(self):
        a = self.full_employee()
        b = self.full_employee()
        self.post("lock")
        self.post("unlock", reason="Month reopened for corrections")
        entries = PayrollAuditLog.objects.filter(action="UNLOCKED")
        self.assertEqual(entries.count(), 2)
        self.assertEqual({e.employee_id for e in entries}, {a.pk, b.pk})
        self.assertEqual({e.reason for e in entries}, {"Month reopened for corrections"})

    def test_unlock_with_nothing_locked_needs_no_reason(self):
        self.full_employee()
        response = self.post("unlock", reason="")
        self.assertTrue(any("None of the employees" in m for m in self.messages(response)))

    def test_relocking_after_unlock_adds_a_second_locked_entry(self):
        a = self.full_employee()
        self.post("lock")
        self.post("unlock", reason="Fixing a wrong attendance entry")
        self.post("lock")
        self.assertEqual(PayrollAuditLog.objects.filter(employee=a, action="LOCKED").count(), 2)

    def test_entries_cannot_be_changed_or_deleted(self):
        self.full_employee()
        self.post("lock")
        entry = PayrollAuditLog.objects.get()
        entry.reason = "edited"
        with self.assertRaises(ValidationError):
            entry.save()
        with self.assertRaises(ValidationError):
            entry.delete()
        self.assertEqual(PayrollAuditLog.objects.get().reason, "")

    def test_record_event_without_a_user(self):
        a = self.full_employee()
        self.post("lock")
        entry = PayrollAuditLog.record_event(self.locked_record(a), "EXPORTED")
        self.assertIsNone(entry.user)
        self.assertEqual(entry.user_name, "")


# ── Bank file: pure functions ───────────────────────────────────────────
class BankFileFunctionTests(SimpleTestCase):
    def record(self, **employee_fields):
        return SimpleNamespace(
            employee=fake_employee(**employee_fields), net_pay=employee_fields.pop("net", D("1000"))
        )

    def split(self, **fields):
        net = fields.pop("net", D("1000"))
        rec = SimpleNamespace(employee=fake_employee(**fields), net_pay=net)
        return split_for_bank([rec])

    def test_complete_details_are_included(self):
        included, left_out = self.split()
        self.assertEqual((len(included), left_out), (1, []))
        self.assertEqual(included[0][1]["account"], "123456789012")
        self.assertEqual(included[0][1]["ref"], "ABCDEF12")

    def test_spaces_are_removed_and_ifsc_is_upper_cased(self):
        included, _ = self.split(bank_account_no="1234 5678 9012", ifsc_no=" hdfc0001234 ")
        self.assertEqual(included[0][1]["account"], "123456789012")
        self.assertEqual(included[0][1]["ifsc"], "HDFC0001234")

    def test_each_problem_leaves_the_employee_out_with_a_reason(self):
        cases = [
            (dict(net=D("0")), "net pay is zero"),
            (dict(bank_account_no=""), "account number is missing"),
            (dict(bank_account_no=None), "account number is missing"),
            (dict(bank_account_no="12-34"), "account number looks invalid"),
            (dict(ifsc_no=""), "IFSC is missing"),
            (dict(ifsc_no="HDFC1001234"), "IFSC looks invalid"),
            (dict(ifsc_no="HDF0001234"), "IFSC looks invalid"),
        ]
        for fields, reason in cases:
            included, left_out = self.split(**fields)
            self.assertEqual(included, [], reason)
            self.assertEqual(left_out, [("Zz Test", reason)], reason)

    def test_csv_layout_and_amounts(self):
        included, _ = self.split()
        out = io.StringIO()
        write_bank_csv(out, included, "Salary Sep 2026")
        rows = list(csv.reader(io.StringIO(out.getvalue())))
        self.assertEqual(rows[0], BANK_HEADER)
        self.assertEqual(
            rows[1],
            ["1", "Zz Test", "123456789012", "HDFC0001234", "HDFC", "1000.00",
             "Salary Sep 2026", "ABCDEF12"],
        )

    def test_text_that_looks_like_a_formula_is_defused(self):
        included, _ = self.split(bank_name="=SUM(1)", user=FakeUser(full="-bad name"))
        out = io.StringIO()
        write_bank_csv(out, included, "Salary Sep 2026")
        row = list(csv.reader(io.StringIO(out.getvalue())))[1]
        self.assertEqual(row[1], "'-bad name")
        self.assertEqual(row[4], "'=SUM(1)")


# ── ECR: pure functions ─────────────────────────────────────────────────
class EcrFunctionTests(SimpleTestCase):
    def record(self, **fields):
        data = dict(
            employee=fake_employee(), pf_wage_base=D("15000"), gross_earned=D("60000"),
            pf=D("1800"), employer_eps=D("1249.50"), employer_epf=D("550.50"),
            days_in_period=30, payable_days=D("30"),
        )
        data.update(fields)
        return SimpleNamespace(**data)

    def test_line_has_eleven_fields_and_contributions_that_add_up(self):
        rec = self.record()
        line = ecr_line(rec, {"uan": "100200300400", "name": "Zz Test"})
        self.assertEqual(
            line,
            "#~#".join(["100200300400", "Zz Test", "60000", "15000", "15000", "15000",
                        "1800", "1250", "550", "0", "0"]),
        )
        fields = line.split("#~#")
        self.assertEqual(len(fields), 11)
        self.assertEqual(int(fields[7]) + int(fields[8]), int(fields[6]))

    def test_ncp_days_are_the_days_not_paid(self):
        rec = self.record(payable_days=D("27.5"), gross_earned=D("55000"))
        fields = ecr_line(rec, {"uan": "100200300400", "name": "Zz Test"}).split("#~#")
        self.assertEqual(fields[9], "3")

    def test_difference_is_never_negative(self):
        rec = self.record(pf=D("1000"))
        fields = ecr_line(rec, {"uan": "100200300400", "name": "Zz Test"}).split("#~#")
        self.assertEqual(fields[8], "0")

    def test_members_without_pf_are_skipped_silently(self):
        rec = self.record(pf=D("0"), employer_eps=D("0"), employer_epf=D("0"))
        self.assertEqual(split_for_ecr([rec]), ([], []))

    def test_uan_problems_leave_the_member_out(self):
        for uan, reason in (("", "UAN is missing"), (None, "UAN is missing"),
                            ("12345678901", "invalid"), ("10020030040A", "invalid")):
            rec = self.record(employee=fake_employee(uan_no=uan))
            included, left_out = split_for_ecr([rec])
            self.assertEqual(included, [], uan)
            self.assertIn(reason, left_out[0][1], uan)

    def test_separator_characters_are_removed_from_names(self):
        rec = self.record(employee=fake_employee(user=FakeUser(full="A#~#B\nC")))
        included, _ = split_for_ecr([rec])
        self.assertEqual(included[0][1]["name"], "A B C")


# ── ESI return: pure functions ──────────────────────────────────────────
class EsiFunctionTests(SimpleTestCase):
    def record(self, **fields):
        data = dict(
            employee=fake_employee(), esi_wage_base=D("15000"), esi=D("112.50"),
            payable_days=D("30"), period_start=date(2026, 9, 1), period_end=date(2026, 9, 30),
        )
        data.update(fields)
        return SimpleNamespace(**data)

    def test_normal_row(self):
        included, left_out = split_for_esi([self.record()], {})
        self.assertEqual(left_out, [])
        d = included[0][1]
        self.assertEqual((d["ip"], d["days"], d["wages"], d["reason"], d["last"]),
                         ("1234567890", 30, 15000, 0, ""))

    def test_half_days_round_up(self):
        included, _ = split_for_esi([self.record(payable_days=D("14.5"))], {})
        self.assertEqual(included[0][1]["days"], 15)

    def test_zero_days_use_reason_two_when_the_employee_left_this_month(self):
        rec = self.record(payable_days=D("0"))
        included, _ = split_for_esi([rec], {1: date(2026, 9, 10)})
        d = included[0][1]
        self.assertEqual((d["days"], d["wages"], d["reason"], d["last"]), (0, 0, 2, "10/09/2026"))

    def test_zero_days_without_an_exit_use_reason_one(self):
        included, _ = split_for_esi([self.record(payable_days=D("0"))], {})
        d = included[0][1]
        self.assertEqual((d["reason"], d["last"]), (1, ""))

    def test_an_exit_in_another_month_is_not_shown(self):
        included, _ = split_for_esi([self.record()], {1: date(2026, 8, 31)})
        self.assertEqual(included[0][1]["last"], "")

    def test_people_without_esi_are_skipped_silently(self):
        rec = self.record(esi_wage_base=D("0"), esi=D("0"))
        self.assertEqual(split_for_esi([rec], {}), ([], []))

    def test_esic_number_problems_leave_the_person_out(self):
        for number, reason in (("", "missing"), (None, "missing"), ("123456789", "invalid"),
                               ("12345678AB", "invalid")):
            rec = self.record(employee=fake_employee(esic_no=number))
            included, left_out = split_for_esi([rec], {})
            self.assertEqual(included, [], number)
            self.assertIn(reason, left_out[0][1], number)

    def test_workbook_header_and_row(self):
        included, _ = split_for_esi([self.record()], {})
        ws = load_workbook(io.BytesIO(build_esi_workbook(included))).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), ESI_HEADER)
        self.assertEqual(rows[1][:5], ("1234567890", "Zz Test", 30, 15000, 0))


# ── Export views ────────────────────────────────────────────────────────
class ExportViewBase(PayrollViewBase):
    def export_url(self, name):
        return reverse(name, args=[YEAR, MONTH])

    def set_esi_profile(self, employee):
        EmployeeStatutoryProfile.objects.update_or_create(
            employee=employee, defaults={"pf_mode": "NA", "esi_mode": "AUTO"}
        )

    def bank_employee(self, **extra):
        fields = dict(bank_account_no="123456789012", ifsc_no="HDFC0001234", bank_name="HDFC")
        fields.update(extra)
        return self.full_employee(**fields)


class ExportAccessTests(ExportViewBase):
    def test_anonymous_users_are_sent_to_login(self):
        self.client.logout()
        for name in ("payrollBankFile", "payrollEcr", "payrollEsiReturn"):
            response = self.client.get(self.export_url(name))
            self.assertEqual(response.status_code, 302, name)
            self.assertIn("login", response.url.lower(), name)

    def test_users_without_the_permission_get_no_file(self):
        nobody = self.make(self.org1, "zz_nobody@example.com")
        self.client.force_login(nobody.user)
        for name in ("payrollBankFile", "payrollEcr", "payrollEsiReturn"):
            response = self.client.get(self.export_url(name))
            self.assertEqual(response.status_code, 302, name)
            self.assertNotIn("attachment", response.get("Content-Disposition", ""), name)

    def test_nothing_locked_means_no_file_and_a_message(self):
        self.bank_employee()
        for name in ("payrollBankFile", "payrollEcr", "payrollEsiReturn"):
            response = self.client.get(self.export_url(name), follow=True)
            self.assertTrue(any("no locked payroll" in m.lower() for m in self.messages(response)), name)
            self.assertNotIn("attachment", response.get("Content-Disposition", ""), name)


class BankFileViewTests(ExportViewBase):
    def test_csv_for_a_locked_employee(self):
        self.bank_employee()
        self.post("lock")
        response = self.client.get(self.export_url("payrollBankFile"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertTrue(response["Content-Disposition"].startswith('attachment; filename="bank_file_'))
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
        self.assertEqual(rows[0], BANK_HEADER)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1:7], ["Zz Test", "123456789012", "HDFC0001234", "HDFC", "58200.00",
                                         "Salary Sep 2026"])

    def test_employees_with_missing_details_are_left_out_and_reported(self):
        self.bank_employee()
        self.full_employee()
        self.post("lock")
        response = self.client.get(self.export_url("payrollBankFile"))
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
        self.assertEqual(len(rows), 2)
        page = self.get()
        self.assertTrue(any("left out of the bank file" in m for m in self.messages(page)))

    def test_no_file_when_nobody_has_complete_details(self):
        self.full_employee()
        self.post("lock")
        response = self.client.get(self.export_url("payrollBankFile"), follow=True)
        self.assertTrue(any("No bank file was made" in m for m in self.messages(response)))
        self.assertNotIn("text/csv", response.get("Content-Type", ""))
        self.assertFalse(PayrollAuditLog.objects.filter(action="EXPORTED").exists())

    def test_the_export_is_logged_per_employee(self):
        a = self.bank_employee()
        self.post("lock")
        self.client.get(self.export_url("payrollBankFile"))
        entry = PayrollAuditLog.objects.get(action="EXPORTED")
        self.assertEqual((entry.employee_id, entry.reason, entry.user_name),
                         (a.pk, "Bank file (NEFT CSV)", HR_USERNAME))


class EcrViewTests(ExportViewBase):
    def lines(self, response):
        return response.content.decode("utf-8").strip().split("\n")

    def test_ecr_line_for_a_locked_employee(self):
        self.bank_employee(uan_no="100200300400")
        self.post("lock")
        response = self.client.get(self.export_url("payrollEcr"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/plain", response["Content-Type"])
        self.assertTrue(response["Content-Disposition"].startswith('attachment; filename="ECR_'))
        lines = self.lines(response)
        self.assertEqual(len(lines), 1)
        self.assertEqual(
            lines[0].split("#~#"),
            ["100200300400", "Zz Test", "60000", "15000", "15000", "15000",
             "1800", "1250", "550", "0", "0"],
        )

    def test_days_not_paid_become_ncp_days(self):
        p = self.employee(uan_no="100200300400")
        self.full_attendance(p, skip={9, 22})
        self.post("lock")
        fields = self.lines(self.client.get(self.export_url("payrollEcr")))[0].split("#~#")
        self.assertEqual(fields[2], "56000")
        self.assertEqual(fields[9], "2")

    def test_member_without_pf_is_not_in_the_file(self):
        self.full_employee(uan_no="100200300400", pf="0")
        self.full_employee(uan_no="100200300401")
        self.post("lock")
        lines = self.lines(self.client.get(self.export_url("payrollEcr")))
        self.assertEqual(len(lines), 1)
        self.assertTrue(lines[0].startswith("100200300401#~#"))

    def test_member_without_a_valid_uan_is_left_out_and_reported(self):
        self.full_employee(uan_no="100200300400")
        self.full_employee(uan_no="12345")
        self.post("lock")
        response = self.client.get(self.export_url("payrollEcr"))
        self.assertEqual(len(self.lines(response)), 1)
        page = self.get()
        self.assertTrue(any("left out of the ECR" in m for m in self.messages(page)))

    def test_no_file_when_nobody_has_a_uan(self):
        self.full_employee()
        self.post("lock")
        response = self.client.get(self.export_url("payrollEcr"), follow=True)
        self.assertTrue(any("No ECR was made" in m for m in self.messages(response)))

    def test_the_export_is_logged(self):
        self.full_employee(uan_no="100200300400")
        self.post("lock")
        self.client.get(self.export_url("payrollEcr"))
        self.assertEqual(PayrollAuditLog.objects.get(action="EXPORTED").reason, "PF ECR")


class EsiReturnViewTests(ExportViewBase):
    def esi_employee(self, **extra):
        fields = dict(esic_no="1234567890")
        fields.update(extra)
        p = self.full_employee(basic="15000", hra="0", pf="0", **fields)
        self.set_esi_profile(p)
        return p

    def sheet(self, response):
        return list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))

    def test_workbook_for_a_locked_employee(self):
        self.esi_employee()
        self.post("lock")
        response = self.client.get(self.export_url("payrollEsiReturn"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertTrue(response["Content-Disposition"].startswith('attachment; filename="ESI_'))
        rows = self.sheet(response)
        self.assertEqual(list(rows[0]), ESI_HEADER)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:5], ("1234567890", "Zz Test", 30, 15000, 0))
        self.assertIn(rows[1][5], (None, ""))

    def test_people_above_the_esi_ceiling_are_not_listed(self):
        self.esi_employee()
        self.full_employee(esic_no="1234567891")
        self.post("lock")
        rows = self.sheet(self.client.get(self.export_url("payrollEsiReturn")))
        self.assertEqual(len(rows), 2)

    def test_missing_esic_number_is_left_out_and_reported(self):
        self.esi_employee()
        self.esi_employee(esic_no="")
        self.post("lock")
        rows = self.sheet(self.client.get(self.export_url("payrollEsiReturn")))
        self.assertEqual(len(rows), 2)
        page = self.get()
        self.assertTrue(any("left out of the ESI return" in m for m in self.messages(page)))

    def test_no_file_when_nobody_qualifies(self):
        self.full_employee()
        self.post("lock")
        response = self.client.get(self.export_url("payrollEsiReturn"), follow=True)
        self.assertTrue(any("No ESI return was made" in m for m in self.messages(response)))

    def test_the_export_is_logged(self):
        self.esi_employee()
        self.post("lock")
        self.client.get(self.export_url("payrollEsiReturn"))
        self.assertEqual(PayrollAuditLog.objects.get(action="EXPORTED").reason, "ESI return")
